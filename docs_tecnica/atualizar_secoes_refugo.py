from docx import Document
from docx.shared import Pt

CAMINHO = "SIAMP_Documentacao_Tecnica.docx"

doc = Document(CAMINHO)


def encontrar(texto_contido):
    for p in doc.paragraphs:
        if texto_contido in p.text:
            return p
    raise RuntimeError(f"Parágrafo de referência não encontrado: {texto_contido!r}")


def estilo_corpo_de(paragrafo_modelo):
    def aplicar(p, texto):
        p.alignment = paragrafo_modelo.alignment
        p.paragraph_format.space_after = paragrafo_modelo.paragraph_format.space_after
        p.paragraph_format.line_spacing = paragrafo_modelo.paragraph_format.line_spacing
        r = p.add_run(texto)
        r.font.size = Pt(12)
        return p
    return aplicar


# --- 1. Seção 5.1: SUPERVISOR não gerencia mais peças/OPs ---------------
p_51 = encontrar("SUPERVISOR (pode cadastrar")
assert len(p_51.runs) == 1
p_51.runs[0].text = (
    "O acesso ao sistema não possui cadastro público: contas são "
    "criadas exclusivamente por um administrador. Três perfis de "
    "usuário controlam o acesso às funcionalidades — ADMIN (acesso "
    "completo, incluindo gestão de usuários, destinatários de "
    "relatório, e cadastro/edição de peças e Ordens de Produção), "
    "SUPERVISOR (cadastra e edita máquinas, além das funcionalidades "
    "de OPERADOR) e OPERADOR (acesso ao apontamento, histórico e "
    "dashboard). Após um pedido direto do setor de planejamento e "
    "controle de produção (PCP) da empresa, o cadastro e a edição de "
    "peças e de Ordens de Produção — antes também acessíveis a "
    "SUPERVISOR — tornaram-se exclusivos de ADMIN; a leitura desses "
    "dois recursos (listagem, usada para popular o seletor de peça e "
    "de Ordem de Produção na tela de apontamento, e o comparativo "
    "exibido no Dashboard) foi deliberadamente mantida sem restrição "
    "de perfil, por ser uma dependência do apontamento diário de "
    "qualquer usuário, independentemente do que pode ou não gerenciar. "
    "A Figura 3 mostra a tela de login, com a identidade visual da "
    "empresa aplicada."
)
print("1. Seção 5.1 corrigida.")


# --- 2. Seção 5.3: ciclo/cavidades informados não têm mais prioridade ---
p_53 = encontrar("os dois valores têm prioridade")
assert len(p_53.runs) == 1
p_53.runs[0].text = (
    "Duas melhorias foram acrescentadas ao cálculo a partir de "
    "feedback direto de quem opera o sistema no chão de fábrica, e "
    "uma delas passou por uma correção de rumo depois de aplicada. A "
    "primeira é a possibilidade de o operador informar manualmente, "
    "no próprio lançamento, tanto o ciclo real observado quanto o "
    "número de cavidades efetivamente em uso na injetora naquele "
    "momento (por exemplo, quando uma cavidade do molde está "
    "temporariamente tamponada). Na primeira versão, os dois valores "
    "tinham prioridade máxima sobre o cadastro da peça no cálculo da "
    "própria capacidade esperada — decisão revertida depois, por "
    "pedido do usuário: o valor \"esperado\" deve permanecer estável "
    "e comparável entre turnos, dependendo só do cadastro oficial da "
    "peça, nunca do que foi digitado num lançamento específico. Ciclo "
    "e cavidades informados continuam sendo registrados e exibidos "
    "lado a lado com o cadastro, destacando divergências, mas "
    "puramente como referência de comparação — nunca mais "
    "influenciam o cálculo em si. O relatório de fechamento de turno "
    "mostra os dois valores lado a lado em cada linha (\"ciclo "
    "cadastrado: 18.5s (informado pelo operador: 20.0s)\", por "
    "exemplo) — sem essa anotação, não havia como diagnosticar, a "
    "partir do relatório sozinho, se uma divergência entre produção "
    "real e esperada vinha de um valor digitado impreciso ou de um "
    "cadastro desatualizado. A Figura 7 mostra os dois valores "
    "informados simultaneamente, com divergência em ambos."
)
print("2. Seção 5.3 (ciclo/cavidades) corrigida.")


# --- 3. Seção 5.3: nova adição sobre Refugo/Índice de Qualidade --------
ref_heading_54 = None
for p in doc.paragraphs:
    if p.text.strip() == "5.4 Salvamento de Rascunho do Turno" and p.style and p.style.name == "Heading 2":
        ref_heading_54 = p
        break
if ref_heading_54 is None:
    raise RuntimeError("Não encontrou a heading '5.4 Salvamento de Rascunho do Turno'.")

ref_corpo = encontrar("A segunda melhoria trata de um problema real")
aplicar = estilo_corpo_de(ref_corpo)

p_refugo1 = ref_heading_54.insert_paragraph_before()
aplicar(
    p_refugo1,
    "Uma terceira funcionalidade foi acrescentada ao cálculo do "
    "indicador de qualidade depois de um pedido do setor de "
    "planejamento e controle de produção da empresa, acompanhado de "
    "um exemplo real de formulário físico já em uso no chão de "
    "fábrica: a possibilidade de registrar o refugo (peças "
    "descartadas) de um lançamento a partir do peso de um lote pesado "
    "numa balança, em vez de exigir contagem manual peça por peça. A "
    "fórmula é simples — a quantidade de refugo é o peso bruto do "
    "lote descartado dividido pelo peso de uma única peça, cadastrado "
    "previamente no catálogo —, mas depende de manter as duas "
    "grandezas na mesma unidade; a primeira implementação presumiu, "
    "incorretamente, que o peso registrado no formulário físico da "
    "empresa estava em quilogramas, aplicando uma conversão "
    "desnecessária antes da divisão. A correção só veio depois que o "
    "próprio usuário confirmou o peso real de uma peça específica "
    "(0,0921 grama, menos de um grama) ao tentar cadastrá-lo e "
    "esbarrar numa validação do formulário que só aceitava valores a "
    "partir de 0,1 — episódio que ilustra o risco de presumir a "
    "unidade de uma grandeza física só pela ordem de grandeza do "
    "número, sem confirmar com quem forneceu o dado original. "
    "Corrigida a unidade (gramas nos dois lados da divisão, sem "
    "conversão), o refugo estimado passou a alimentar de fato o "
    "Índice de Qualidade e o OEE do turno — antes dessa "
    "funcionalidade, o modelo de lançamentos livres sempre assumia "
    "100% de qualidade, por não ter nenhuma fonte de dado sobre "
    "peças ruins."
)

p_refugo2 = ref_heading_54.insert_paragraph_before()
aplicar(
    p_refugo2,
    "Um erro de agregação foi identificado já em uso real, comparando "
    "dois relatórios do mesmo turno: o total de \"Peças Boas\" exibido "
    "no relatório de fechamento correspondia exatamente à produção de "
    "uma única injetora — a única, entre as sete em produção naquele "
    "turno, que teve descarte pesado —, como se as demais não "
    "tivessem produzido nada. A causa estava na forma como o total "
    "era somado: cada lançamento só contribuía para o total de peças "
    "boas quando ele próprio tinha um refugo calculável, deixando de "
    "fora inteiramente os lançamentos sem descarte pesado, a grande "
    "maioria na prática. A correção redefiniu o total de peças boas "
    "como a produção do turno inteiro menos o refugo total do turno, "
    "garantindo que os dois números sempre somem exatamente a "
    "produção total — identidade que o relatório agora respeita em "
    "qualquer combinação de injetoras com ou sem descarte pesado "
    "naquele turno. Um segundo bug, também relatado pelo próprio "
    "usuário, foi a perda silenciosa do peso do descarte — e do ciclo "
    "e das cavidades informados — sempre que um lançamento de um "
    "turno já fechado era reeditado: a tela de edição reconstruía o "
    "estado de cada lançamento a partir de uma lista de campos "
    "definida manualmente, que não havia sido atualizada quando os "
    "campos novos foram introduzidos, então bastava abrir a tela de "
    "edição para os dados desaparecerem, mesmo sem qualquer alteração "
    "deliberada por parte do usuário."
)
print("3. Nova seção sobre Refugo adicionada (2 parágrafos).")


# --- 4. Seção 5.5: peças não são mais gerenciadas por supervisor -------
p_55 = encontrar("Administradores e supervisores cadastram e mantêm o parque")
assert len(p_55.runs) == 1
p_55.runs[0].text = (
    "Administradores e supervisores cadastram e mantêm o parque de "
    "injetoras diretamente pela interface, sem qualquer valor fixo no "
    "código-fonte da aplicação; o catálogo de peças segue a mesma "
    "lógica de cadastro livre, mas seu cadastro e sua edição "
    "tornaram-se exclusivos do perfil ADMIN (seção 5.1) — qualquer "
    "perfil continua podendo selecionar uma peça já cadastrada em "
    "qualquer tela de apontamento. Uma revisão importante de "
    "modelagem tornou o ciclo médio e o número de cavidades campos "
    "obrigatórios no cadastro da peça (antes opcionais, com a máquina "
    "servindo de valor padrão) — decisão que reflete a realidade de "
    "que uma mesma injetora frequentemente troca de molde entre "
    "turnos, tornando o ciclo da máquina, isoladamente, uma "
    "referência pouco confiável para o cálculo de OEE. O cadastro "
    "também aceita, opcionalmente, o peso de uma unidade da peça em "
    "gramas, usado para estimar a quantidade de refugo a partir do "
    "peso de um lote descartado (seção 5.3). Esses campos continuam "
    "editáveis após o cadastro, contemplando situações como o "
    "fechamento temporário de uma cavidade do molde. Para catálogos "
    "maiores, a seleção de peça no apontamento conta com um campo de "
    "busca por código ou descrição, evitando a necessidade de rolar "
    "uma lista longa. A Figura 8 mostra a tela de cadastro de peças."
)
print("4. Seção 5.5 corrigida.")


# --- 5. Seção 7: contagem de testes e descrição desatualizada ----------
p_7 = encontrar("totalizando 277 testes")
assert len(p_7.runs) == 1
texto7 = p_7.runs[0].text
texto7 = texto7.replace("totalizando 277 testes", "totalizando 301 testes")
texto7 = texto7.replace(
    "prevalência do ciclo e das cavidades informados manualmente sobre "
    "o cadastro da peça,",
    "uso exclusivo do cadastro da peça no cálculo mesmo quando ciclo e "
    "cavidades são informados manualmente,",
)
texto7 = texto7.replace(
    "e os dois achados de segurança descritos na seção 6.",
    "os dois achados de segurança descritos na seção 6, o cálculo de "
    "refugo por peso de descarte (incluindo a agregação correta de "
    "\"Peças Boas\" para o turno inteiro) e a restrição de perfil "
    "para cadastro de peças e Ordens de Produção descrita na seção "
    "5.1.",
)
p_7.runs[0].text = texto7
print("5. Seção 7 (testes) corrigida.")

doc.save(CAMINHO)
print("\nDocumento salvo com sucesso.")
